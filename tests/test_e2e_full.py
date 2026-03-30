"""Test E2E complet de toutes les fonctionnalités de l'app ICAE Streamlit.

Modules testés :
  1. ICAE : import → pipeline → export (write_icae_output)
  2. Prévisions : séries HF → forecast → injection → recalc export (write_icae_recalc_output)
  3. Nowcast : PIB + HF trimestriels → run_nowcast → alignment info → export (write_nowcast_excel)
  4. CEMAC : multi-pays → agrégation → export (write_cemac_excel)
"""

import sys, os, io, traceback
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import numpy as np
import pandas as pd
import openpyxl

# ─── Paths ─────────────────────────────────────────────────────────────────
from config import COUNTRY_CODES, COUNTRY_NAMES, PIB_2014, POIDS_PIB, CEMAC_TEMPLATE
from pathlib import Path

LIVRABLE = Path(r"c:\Users\HP\Documents\Stage pro BEAC\Work\ICAE\Mars 2026\Livrable_Final\01_Classeurs_ICAE")
OUTPUT_DIR = Path(os.path.dirname(os.path.abspath(__file__))) / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

# Test files — use "_now" versions if available
def _find_country_file(code):
    for suffix in ("_Consolide_now.xlsx", "_Consolide.xlsx"):
        p = LIVRABLE / f"ICAE_{code}{suffix}"
        if p.exists():
            return p
    return None

# ─── Helpers ───────────────────────────────────────────────────────────────
PASS = 0
FAIL = 0
ERRORS = []

def check(name, condition, detail=""):
    global PASS, FAIL
    if condition:
        PASS += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        msg = f"  ❌ {name}" + (f" — {detail}" if detail else "")
        print(msg)
        ERRORS.append(msg)


def section(title):
    print(f"\n{'='*70}\n  {title}\n{'='*70}")


# ═══════════════════════════════════════════════════════════════════════════
#  MODULE 1 — ICAE : Import → Pipeline → Export
# ═══════════════════════════════════════════════════════════════════════════
section("MODULE 1 — ICAE (Import → Pipeline → Export)")

from io_utils.excel_reader import (
    list_sheets, read_consignes, read_codification,
    read_donnees_calcul, rename_columns_to_codes,
)
from core.icae_engine import run_icae_pipeline
from core.quarterly import quarterly_mean, calc_ga_trim, calc_gt_trim
from io_utils.excel_writer import write_icae_output

TEST_CODE = "CMR"
test_file = _find_country_file(TEST_CODE)
check("Fichier test CMR trouvé", test_file is not None and test_file.exists(),
      f"Cherché dans {LIVRABLE}")

if test_file:
    # 1a. Import
    sheets = list_sheets(test_file)
    check("Feuilles attendues présentes",
          all(s in sheets for s in ["Consignes", "Codification", "Donnees_calcul", "CALCUL_ICAE"]),
          f"Trouvées : {sheets}")

    consignes = read_consignes(test_file)
    check("Consignes lues", "base_year" in consignes, f"Consignes = {consignes}")

    codif = read_codification(test_file)
    check("Codification lue", len(codif) > 0, f"{len(codif)} variables")

    donnees = read_donnees_calcul(test_file)
    donnees = rename_columns_to_codes(donnees, codif)
    check("Donnees_calcul lues", len(donnees) > 50 and "Date" in donnees.columns,
          f"{len(donnees)} lignes, {len(donnees.columns)} cols")

    # 1b. Pipeline ICAE
    base_year = consignes.get("base_year", 2023)
    dates = pd.to_datetime(donnees["Date"])
    base_mask = dates.dt.year == base_year
    base_indices = donnees.index[base_mask]
    if len(base_indices) > 0:
        base_rows = range(base_indices[0], base_indices[-1] + 1)
    else:
        base_rows = range(108, 120)

    priors = pd.Series(dtype=float)
    if "Code" in codif.columns and "PRIOR" in codif.columns:
        priors = pd.Series(codif["PRIOR"].values, index=codif["Code"].values, dtype=float).fillna(0)
    else:
        data_cols = [c for c in donnees.columns if c != "Date"]
        priors = pd.Series(1.0, index=data_cols)

    results = run_icae_pipeline(donnees=donnees, priors=priors,
                                base_year=base_year, base_rows=base_rows)

    check("Pipeline ICAE exécuté", "icae" in results and "dates" in results)
    icae = results["icae"]
    check("ICAE calculé", len(icae) > 100, f"{len(icae)} valeurs")
    check("ICAE dans plage [50-200]",
          50 <= icae.dropna().median() <= 200,
          f"Médiane = {icae.dropna().median():.2f}")

    # Quarterly
    q = quarterly_mean(icae, dates)
    check("Trimestriel calculé", len(q) > 10, f"{len(q)} trimestres")

    # 1c. Export — write_icae_output
    results["quarterly"] = q
    results["dates"] = dates
    out_bytes = write_icae_output(test_file, results, TEST_CODE)
    check("Export write_icae_output", len(out_bytes) > 1000)

    # Vérifier le classeur exporté
    wb_out = openpyxl.load_workbook(io.BytesIO(out_bytes))
    check("Feuilles dans l'export",
          "CALCUL_ICAE" in wb_out.sheetnames and "Resultats_Trim" in wb_out.sheetnames,
          f"Feuilles = {wb_out.sheetnames}")

    # Vérifier Resultats_Trim
    ws_rt = wb_out["Resultats_Trim"]
    rt_r1 = ws_rt.cell(1, 1).value
    rt_r2_b = ws_rt.cell(2, 2).value
    check("Resultats_Trim a un header",
          rt_r1 is not None and "Trim" in str(rt_r1),
          f"A1 = {rt_r1}")
    # Trouver la dernière ligne non-vide
    rt_last = 1
    for r in range(ws_rt.max_row, 1, -1):
        if ws_rt.cell(r, 1).value is not None:
            rt_last = r
            break
    n_trims_out = rt_last - 1
    expected_trims = len(q)
    check("Resultats_Trim — nombre de trimestres",
          n_trims_out == expected_trims,
          f"Attendu={expected_trims}, trouvé={n_trims_out}")

    # Vérifier formules dans Resultats_Trim
    sample_formula = str(ws_rt.cell(10, 2).value or "")
    check("Resultats_Trim — formule AVERAGE",
          "AVERAGE" in sample_formula or "CALCUL_ICAE" in sample_formula,
          f"B10 = {sample_formula}")
    if n_trims_out > 5:
        ga_formula = str(ws_rt.cell(7, 3).value or "")
        check("Resultats_Trim — formule GA (row 7)",
              "B7" in ga_formula or "B3" in ga_formula or ga_formula.startswith("="),
              f"C7 = {ga_formula}")
    wb_out.close()

    # Date range
    first_date = dates.iloc[0]
    last_date = dates.iloc[-1]
    print(f"  ℹ️  Période données : {first_date.strftime('%Y-%m')} → {last_date.strftime('%Y-%m')} ({len(donnees)} mois)")
    print(f"  ℹ️  ICAE médiane = {icae.dropna().median():.2f}, min={icae.dropna().min():.2f}, max={icae.dropna().max():.2f}")


# ═══════════════════════════════════════════════════════════════════════════
#  MODULE 2 — PRÉVISIONS : Forecast → Injection → Recalc Export
# ═══════════════════════════════════════════════════════════════════════════
section("MODULE 2 — PRÉVISIONS (Forecast → Injection → Recalc Export)")

from core.forecast_engine import run_all_forecasts, get_methods_for_frequency
from io_utils.excel_writer import write_previsions_excel, write_icae_recalc_output

if test_file:
    # Préparer les données : prendre 3 séries actives
    active_codes = []
    if "Code" in codif.columns and "PRIOR" in codif.columns:
        for _, row in codif.iterrows():
            try:
                if float(row.get("PRIOR", 0)) > 0 and str(row.get("Statut", "Actif")).lower() != "inactif":
                    active_codes.append(str(row["Code"]).strip())
            except (ValueError, TypeError):
                pass
    if not active_codes:
        active_codes = [c for c in donnees.columns if c != "Date"][:3]
    test_vars = [c for c in active_codes[:3] if c in donnees.columns]
    check("Variables actives pour prévision", len(test_vars) > 0, f"Vars = {test_vars}")

    horizon = 3  # 3 mois de prévision
    all_forecasts = {}
    all_backtesting = {}
    all_recommandations = {}

    for var in test_vars:
        series = pd.to_numeric(donnees[var], errors="coerce").dropna()
        if len(series) < 24:
            continue
        methods = get_methods_for_frequency("Mensuelle")
        res = run_all_forecasts(series, horizon, methods)
        all_forecasts[var] = res.get("forecasts", {})
        all_backtesting[var] = res.get("backtesting", {})
        best = res.get("best_method", methods[0])
        all_recommandations[var] = {
            "method": best,
            "mape": res.get("backtesting", {}).get(best, {}).get("mape", np.nan),
        }

    check("Prévisions calculées", len(all_forecasts) > 0, f"{len(all_forecasts)} variables")

    # 2a. Export prévisions brutes (write_previsions_excel)
    prev_bytes = write_previsions_excel(
        donnees, all_forecasts, all_backtesting, all_recommandations
    )
    check("Export write_previsions_excel", len(prev_bytes) > 1000)
    wb_prev = openpyxl.load_workbook(io.BytesIO(prev_bytes))
    expected_sheets_prev = ["Donnees_Completes", "Donnees", "Previsions", "Backtesting", "Recommandations"]
    for s in expected_sheets_prev:
        check(f"Feuille '{s}' dans export prévisions", s in wb_prev.sheetnames)
    wb_prev.close()

    # 2b. Injection des prévisions → recalc ICAE
    # Simuler l'injection : prolonger les données avec les prévisions
    last_date = pd.to_datetime(donnees["Date"]).max()
    fc_dates = pd.date_range(last_date + pd.DateOffset(months=1), periods=horizon, freq="MS")
    new_rows = pd.DataFrame({"Date": fc_dates})
    for var in test_vars:
        best = all_recommandations[var]["method"]
        fc = all_forecasts[var].get(best, np.full(horizon, np.nan))
        new_rows[var] = fc[:horizon] if len(fc) >= horizon else np.concatenate(
            [fc, np.full(horizon - len(fc), np.nan)])
    # Remplir les autres colonnes avec NaN
    for col in donnees.columns:
        if col not in new_rows.columns:
            new_rows[col] = np.nan

    extended = pd.concat([donnees, new_rows], ignore_index=True)
    check("Données prolongées",
          len(extended) == len(donnees) + horizon,
          f"{len(donnees)} + {horizon} = {len(extended)}")

    # Recalcul ICAE sur les données prolongées
    ext_dates = pd.to_datetime(extended["Date"])
    ext_base_mask = ext_dates.dt.year == base_year
    ext_base_indices = extended.index[ext_base_mask]
    if len(ext_base_indices) > 0:
        ext_base_rows = range(ext_base_indices[0], ext_base_indices[-1] + 1)
    else:
        ext_base_rows = base_rows

    results_ext = run_icae_pipeline(donnees=extended, priors=priors,
                                     base_year=base_year, base_rows=ext_base_rows)
    icae_ext = results_ext["icae"]
    check("ICAE recalculé (étendu)",
          len(icae_ext) == len(extended),
          f"{len(icae_ext)} valeurs (attendu {len(extended)})")

    # Quarterly étendu
    q_ext = quarterly_mean(icae_ext, ext_dates)
    check("Trimestriel étendu", len(q_ext) > len(q))
    last_q_label = str(q_ext.iloc[-1].get("trimestre", q_ext.iloc[-1].get("quarter", "")))
    print(f"  ℹ️  Dernier trimestre étendu : {last_q_label}")

    # 2c. Export recalc (write_icae_recalc_output) — avec template
    results_ext["dates"] = ext_dates
    recalc_bytes = write_icae_recalc_output(
        donnees=extended, results=results_ext,
        quarterly=q_ext, contrib_trim=None,
        codification=codif, country_code=TEST_CODE,
        source_path=test_file,
    )
    check("Export write_icae_recalc_output (template)", len(recalc_bytes) > 5000)

    # Vérifier le classeur recalculé
    wb_recalc = openpyxl.load_workbook(io.BytesIO(recalc_bytes))
    check("Feuilles dans recalc export",
          all(s in wb_recalc.sheetnames for s in
              ["Donnees_calcul", "CALCUL_ICAE", "Resultats_Trim", "Contrib"]),
          f"Feuilles = {wb_recalc.sheetnames}")

    # Vérifier Donnees_calcul étendu
    ws_d = wb_recalc["Donnees_calcul"]
    from io_utils.excel_writer import _actual_last_row
    d_last = _actual_last_row(ws_d)
    expected_d_rows = len(extended) + 1  # header + data
    check("Donnees_calcul — lignes étendues",
          d_last == expected_d_rows,
          f"Attendu={expected_d_rows}, trouvé={d_last}")
    # Dernière date
    last_date_val = ws_d.cell(d_last, 1).value
    check("Donnees_calcul — dernière date",
          last_date_val is not None,
          f"A{d_last} = {last_date_val}")

    # Vérifier CALCUL_ICAE étendu
    ws_c = wb_recalc["CALCUL_ICAE"]
    c_last = _actual_last_row(ws_c)
    # Les lignes ajoutées doivent contenir des formules IFERROR(...)
    sample_tcs = str(ws_c.cell(c_last, 2).value or "")
    check("CALCUL_ICAE — formule TCS dans dernière ligne",
          "IFERROR" in sample_tcs or "Donnees_calcul" in sample_tcs,
          f"B{c_last} = {sample_tcs[:80]}")

    # Vérifier Resultats_Trim étendu
    ws_rt = wb_recalc["Resultats_Trim"]
    rt_last = 1
    for r in range(ws_rt.max_row, 1, -1):
        if ws_rt.cell(r, 1).value is not None:
            rt_last = r
            break
    n_trims_ext = rt_last - 1
    expected_trims_ext = len(q_ext)
    check("Resultats_Trim — trimestres étendus",
          n_trims_ext == expected_trims_ext,
          f"Attendu={expected_trims_ext}, trouvé={n_trims_ext}")
    last_trim_label = str(ws_rt.cell(rt_last, 1).value or "")
    check("Resultats_Trim — dernier trimestre = attendu",
          last_trim_label == last_q_label,
          f"Attendu={last_q_label}, trouvé={last_trim_label}")
    # Formule AVERAGE dans le dernier trimestre
    last_avg = str(ws_rt.cell(rt_last, 2).value or "")
    check("Resultats_Trim — formule AVERAGE dans dernier trim",
          "AVERAGE" in last_avg,
          f"B{rt_last} = {last_avg[:60]}")

    # Vérifier Contrib étendu
    ws_ct = wb_recalc["Contrib"]
    ct_last = _actual_last_row(ws_ct)
    ct_a_last = str(ws_ct.cell(ct_last, 1).value or "")
    check("Contrib — dernière ligne contient une formule",
          "CALCUL_ICAE" in ct_a_last,
          f"A{ct_last} = {ct_a_last[:60]}")

    wb_recalc.close()
    # Sauvegarder pour inspection manuelle
    out_recalc_path = OUTPUT_DIR / f"E2E_ICAE_{TEST_CODE}_recalc.xlsx"
    out_recalc_path.write_bytes(recalc_bytes)
    print(f"  ℹ️  Classeur recalculé sauvegardé : {out_recalc_path}")


# ═══════════════════════════════════════════════════════════════════════════
#  MODULE 3 — NOWCAST : PIB + HF → Modèles → Alignment → Export
# ═══════════════════════════════════════════════════════════════════════════
section("MODULE 3 — NOWCAST (PIB + HF → Modèles → Alignment → Export)")

from core.nowcast_engine import run_nowcast, _prepare_data, _idx_to_qkey
from core.quarterly import agg_m_to_q
from io_utils.excel_writer import write_nowcast_excel

# Construire des données synthétiques PIB annuel + HF mensuels
np.random.seed(42)
n_years = 15
years = list(range(2008, 2008 + n_years))
pib_annual = pd.Series(
    np.cumsum(np.random.randn(n_years) * 50 + 100) + 5000,
    index=years, name="PIB"
)

# HF mensuel : 3 variables couvrant une période plus large que le PIB
n_months_hf = (n_years + 2) * 12  # 2 ans de plus
dates_hf = pd.date_range("2007-01-01", periods=n_months_hf, freq="MS")
hf_monthly = pd.DataFrame({
    "HF1": np.cumsum(np.random.randn(n_months_hf) * 0.5) + 100,
    "HF2": np.cumsum(np.random.randn(n_months_hf) * 0.3) + 200,
    "HF3": np.sin(np.arange(n_months_hf) / 6) * 10 + 50,
}, index=dates_hf)

# Trimestialiser le PIB (uniforme)
from core.tempdisagg import disaggregate_annual_to_quarterly
pib_q = disaggregate_annual_to_quarterly(pib_annual)
check("PIB trimestrialisé", len(pib_q) == n_years * 4, f"{len(pib_q)} trimestres")
print(f"  ℹ️  PIB : {pib_q.index[0]} → {pib_q.index[-1]}")

# Agréger HF en trimestriel
agg_map = {"HF1": "flow", "HF2": "mean", "HF3": "mean"}
hf_q = agg_m_to_q(hf_monthly, dates_hf, agg_map)
if "quarter" in hf_q.columns:
    hf_q = hf_q.set_index("quarter")
check("HF trimestriel", len(hf_q) > 40, f"{len(hf_q)} trimestres")
print(f"  ℹ️  HF : {hf_q.index[0]} → {hf_q.index[-1]}")

# 3a. Tester _prepare_data directement — vérifier _align_info
prep = _prepare_data(pib_q, hf_q)
check("_prepare_data retourne un dict", isinstance(prep, dict))
check("_align_info présent", "_align_info" in prep)
if isinstance(prep, dict) and "_align_info" in prep:
    ai = prep["_align_info"]
    check("_align_info a pib_range", "pib_range" in ai, f"{ai}")
    check("_align_info a hf_range", "hf_range" in ai)
    check("_align_info a common_range", "common_range" in ai)
    check("_align_info n_hf_beyond_pib >= 0",
          ai.get("n_hf_beyond_pib", -1) >= 0,
          f"n_hf_beyond = {ai.get('n_hf_beyond_pib')}")
    print(f"  ℹ️  PIB range : {ai['pib_range']}")
    print(f"  ℹ️  HF range  : {ai['hf_range']}")
    print(f"  ℹ️  Common    : {ai['common_range']}")
    print(f"  ℹ️  HF beyond PIB : {ai['n_hf_beyond_pib']} trimestres")
    print(f"  ℹ️  PIB before HF : {ai['n_pib_before_hf']} trimestres")

# Vérifier que hf_future / index_future ne sont PAS dans le résultat
check("Pas de hf_future dans _prepare_data",
      "hf_future" not in prep,
      "hf_future a été correctement supprimé" if "hf_future" not in prep
      else "⚠️ hf_future est encore présent !")
check("Pas de index_future dans _prepare_data",
      "index_future" not in prep)

# 3b. Exécuter run_nowcast
nw_results = run_nowcast(pib_q, hf_q, models=["Bridge", "U-MIDAS", "PC", "DFM"],
                         h_ahead=4, n_components=2)

check("run_nowcast retourne des résultats",
      any(k in nw_results for k in ["Bridge", "U-MIDAS", "PC", "DFM"]))
check("_align dans run_nowcast", "_align" in nw_results)

if "_align" in nw_results:
    align = nw_results.pop("_align")
    print(f"  ℹ️  Align info via run_nowcast : {align}")

errors = nw_results.pop("_errors", [])
if errors:
    print(f"  ⚠️  Erreurs nowcast : {errors}")

for model_name in ["Bridge", "U-MIDAS", "PC", "DFM"]:
    if model_name in nw_results:
        fc = nw_results[model_name]["forecast"]
        n_valid = fc.dropna().shape[0]
        check(f"Nowcast {model_name} — {n_valid} prévisions valides",
              n_valid > 10,
              f"forecast len = {len(fc)}, non-NaN = {n_valid}")
        # Vérifier que le forecast ne dépasse PAS la période PIB (pas d'extrapolation)
        pib_last = pib_q.index[-1]
        fc_last = fc.dropna().index[-1]
        check(f"Nowcast {model_name} — pas d'extrapolation au-delà du PIB",
              str(fc_last) <= str(pib_last),
              f"PIB max = {pib_last}, Forecast max = {fc_last}")

# 3c. Export Nowcast
params = {"Pays": "SYN", "Modeles": "Bridge, U-MIDAS, PC, DFM",
          "Nb facteurs": 2, "Horizon": 4}
hf_vars = ["HF1", "HF2", "HF3"]
nw_export_bytes = write_nowcast_excel(
    pib_q, nw_results, params,
    hf_vars=hf_vars, agg_map=agg_map,
    hf_df=hf_monthly.reset_index().rename(columns={"index": "Date"})
)
check("Export write_nowcast_excel", len(nw_export_bytes) > 1000)

wb_nw = openpyxl.load_workbook(io.BytesIO(nw_export_bytes))
expected_sheets_nw = ["PIB_and_Nowcasts", "Performance", "Glissement_Annuel",
                      "Indicateurs_HF", "Donnees_HF_mensuelles", "Parametres"]
for s in expected_sheets_nw:
    check(f"Feuille '{s}' dans export Nowcast", s in wb_nw.sheetnames,
          f"Feuilles = {wb_nw.sheetnames}")

# Vérifier contenu PIB_and_Nowcasts
ws_pn = wb_nw["PIB_and_Nowcasts"]
# Compter les colonnes (PIB_observe + 4 modèles = 5 data cols)
pn_headers = [ws_pn.cell(1, c).value for c in range(1, ws_pn.max_column + 1)]
check("PIB_and_Nowcasts — colonnes modèles",
      "PIB_observe" in str(pn_headers),
      f"Headers = {pn_headers}")
# Nombre de lignes de données
pn_last = _actual_last_row(ws_pn)
check("PIB_and_Nowcasts — lignes de données",
      pn_last > 10, f"Dernière ligne = {pn_last}")

# Vérifier Performance
ws_perf = wb_nw["Performance"]
perf_last = _actual_last_row(ws_perf)
check("Performance — 4 lignes de modèles",
      perf_last >= 5,  # header + 4 models
      f"Dernière ligne = {perf_last}")

# Vérifier Indicateurs_HF
ws_hf = wb_nw["Indicateurs_HF"]
hf_last = _actual_last_row(ws_hf)
check("Indicateurs_HF — 3 variables listées",
      hf_last >= 4,  # header + 3
      f"Dernière ligne = {hf_last}")

wb_nw.close()
out_nw_path = OUTPUT_DIR / "E2E_NOWCAST_SYN.xlsx"
out_nw_path.write_bytes(nw_export_bytes)
print(f"  ℹ️  Export Nowcast sauvegardé : {out_nw_path}")


# ═══════════════════════════════════════════════════════════════════════════
#  MODULE 3b — NOWCAST avec fichier réel (PIB depuis classeur)
# ═══════════════════════════════════════════════════════════════════════════
section("MODULE 3b — NOWCAST avec fichier réel (HF depuis classeur CMR)")

if test_file and len(donnees) > 50:
    # HF = séries du classeur CMR
    real_hf_vars = [c for c in test_vars if c in donnees.columns][:5]
    if not real_hf_vars:
        real_hf_vars = [c for c in donnees.columns if c != "Date"][:5]

    hf_real = donnees[["Date"] + real_hf_vars].copy()
    hf_real["Date"] = pd.to_datetime(hf_real["Date"], errors="coerce")
    hf_real = hf_real.dropna(subset=["Date"])
    hf_real_num = hf_real[real_hf_vars].apply(pd.to_numeric, errors="coerce")
    real_dates = hf_real["Date"]

    real_agg = {v: "flow" for v in real_hf_vars}
    hf_q_real = agg_m_to_q(hf_real_num, real_dates, real_agg)
    if "quarter" in hf_q_real.columns:
        hf_q_real = hf_q_real.set_index("quarter")

    check("HF réel agrégé en trimestriel",
          len(hf_q_real) > 10, f"{len(hf_q_real)} trimestres")

    # PIB synthétique couvrant une sous-période
    n_real_q = len(hf_q_real)
    pib_real_q = pd.Series(
        np.cumsum(np.random.randn(n_real_q) * 20 + 50) + 3000,
        index=hf_q_real.index
    )
    # Tronquer le PIB pour simuler un gap : PIB s'arrête 2 trimestres avant les HF
    pib_short = pib_real_q.iloc[:-2]
    print(f"  ℹ️  PIB : {pib_short.index[0]} → {pib_short.index[-1]} ({len(pib_short)} Q)")
    print(f"  ℹ️  HF  : {hf_q_real.index[0]} → {hf_q_real.index[-1]} ({len(hf_q_real)} Q)")

    nw_real = run_nowcast(pib_short, hf_q_real,
                          models=["Bridge", "U-MIDAS", "PC", "DFM"])
    real_align = nw_real.pop("_align", None)
    real_errors = nw_real.pop("_errors", [])

    check("_align info pour données réelles", real_align is not None)
    if real_align:
        check("n_hf_beyond_pib > 0 (HF déborde)",
              real_align.get("n_hf_beyond_pib", 0) > 0,
              f"n_hf_beyond = {real_align.get('n_hf_beyond_pib')}")
        print(f"  ℹ️  Align: PIB={real_align['pib_range']}, HF={real_align['hf_range']}, "
              f"Common={real_align['common_range']}")

    # Vérifier aucun modèle ne dépasse le dernier PIB
    for mn, r in nw_real.items():
        if isinstance(r, dict) and "forecast" in r:
            fc = r["forecast"].dropna()
            if len(fc) > 0:
                fc_max = str(fc.index[-1])
                pib_max = str(pib_short.index[-1])
                check(f"Nowcast réel {mn} — pas d'extrapolation",
                      fc_max <= pib_max,
                      f"Forecast max={fc_max}, PIB max={pib_max}")


# ═══════════════════════════════════════════════════════════════════════════
#  MODULE 4 — CEMAC : Multi-pays → Agrégation → Export
# ═══════════════════════════════════════════════════════════════════════════
section("MODULE 4 — CEMAC (Multi-pays → Agrégation → Export)")

from core.cemac_engine import compute_icae_cemac, quarterly_cemac
from io_utils.excel_writer import write_cemac_excel

# Charger autant de pays que possible
icae_dict = {}
dates_dict = {}
loaded_countries = []

for code in COUNTRY_CODES:
    fpath = _find_country_file(code)
    if fpath is None:
        continue
    try:
        _sh = list_sheets(fpath)
        if "Donnees_calcul" not in _sh:
            continue
        _cons = read_consignes(fpath) if "Consignes" in _sh else {"base_year": 2023}
        _cod = read_codification(fpath) if "Codification" in _sh else pd.DataFrame()
        _don = read_donnees_calcul(fpath)
        if not _cod.empty:
            _don = rename_columns_to_codes(_don, _cod)
        _by = _cons.get("base_year", 2023)
        _dt = pd.to_datetime(_don["Date"])
        _bm = _dt.dt.year == _by
        _bi = _don.index[_bm]
        _br = range(_bi[0], _bi[-1] + 1) if len(_bi) > 0 else range(108, 120)

        _pr = pd.Series(dtype=float)
        if not _cod.empty and "Code" in _cod.columns and "PRIOR" in _cod.columns:
            _pr = pd.Series(_cod["PRIOR"].values, index=_cod["Code"].values, dtype=float).fillna(0)
        else:
            _dc = [c for c in _don.columns if c != "Date"]
            _pr = pd.Series(1.0, index=_dc)

        _res = run_icae_pipeline(donnees=_don, priors=_pr, base_year=_by, base_rows=_br)
        _icae = _res["icae"]
        _icae.index = _dt
        icae_dict[code] = _icae
        dates_dict[code] = _dt
        loaded_countries.append(code)
    except Exception as e:
        print(f"  ⚠️  Erreur chargement {code}: {e}")

check("Pays chargés", len(loaded_countries) >= 2,
      f"{len(loaded_countries)} pays : {loaded_countries}")

if len(loaded_countries) >= 2:
    # Vérifier les plages temporelles
    for code in loaded_countries:
        s = icae_dict[code]
        idx = pd.to_datetime(s.index)
        valid = idx[s.notna().values]
        if len(valid) > 0:
            print(f"  ℹ️  {code}: {valid.min().strftime('%Y-%m')} → {valid.max().strftime('%Y-%m')} "
                  f"({len(valid)} mois)")

    # 4a. Calcul ICAE CEMAC
    result_df = compute_icae_cemac(icae_dict, POIDS_PIB)
    check("ICAE CEMAC calculé", not result_df.empty,
          f"{len(result_df)} lignes × {len(result_df.columns)} cols")
    check("Colonnes ICAE_CEMAC, GA, GT présentes",
          "ICAE_CEMAC" in result_df.columns and "GA" in result_df.columns)

    icae_cemac = result_df["ICAE_CEMAC"]
    check("ICAE CEMAC médiane dans [50-200]",
          50 <= icae_cemac.dropna().median() <= 200,
          f"Médiane = {icae_cemac.dropna().median():.2f}")

    # 4b. Trimestriel
    q_cemac = quarterly_cemac(result_df, pd.to_datetime(result_df.index))
    check("ICAE CEMAC trimestriel", len(q_cemac) > 10,
          f"{len(q_cemac)} trimestres")
    check("Feuille Trimestre + GA_Trim",
          "Trimestre" in q_cemac.columns and "GA_Trim" in q_cemac.columns,
          f"Colonnes = {list(q_cemac.columns)}")

    # 4c. Export CEMAC
    cemac_template = CEMAC_TEMPLATE if Path(CEMAC_TEMPLATE).exists() else None
    cemac_bytes = write_cemac_excel(result_df, q_cemac, POIDS_PIB,
                                    template_path=cemac_template)
    check("Export write_cemac_excel", len(cemac_bytes) > 1000)

    wb_cemac = openpyxl.load_workbook(io.BytesIO(cemac_bytes))
    if cemac_template:
        expected_cemac_sheets = ["Poids_PIB", "ICAE_Pays", "ICAE_Trimestriel"]
        for s in expected_cemac_sheets:
            check(f"Feuille '{s}' dans export CEMAC", s in wb_cemac.sheetnames,
                  f"Feuilles = {wb_cemac.sheetnames}")

        # Vérifier ICAE_Pays
        if "ICAE_Pays" in wb_cemac.sheetnames:
            ws_ip = wb_cemac["ICAE_Pays"]
            ip_last = _actual_last_row(ws_ip)
            expected_ip_rows = 5 + len(result_df) - 1  # data_start=5, 0-indexed
            check("ICAE_Pays — nombre de lignes",
                  abs(ip_last - expected_ip_rows) <= 1,
                  f"Attendu≈{expected_ip_rows}, trouvé={ip_last}")
            # Première date
            first_dt = ws_ip.cell(5, 1).value
            check("ICAE_Pays — A5 est une date",
                  first_dt is not None,
                  f"A5 = {first_dt}")
            # Dernière ligne — vérifier formule ICAE CEMAC
            icae_formula = str(ws_ip.cell(ip_last, 15).value or "")  # Col O
            check("ICAE_Pays — formule CEMAC dans dernière ligne",
                  "IFERROR" in icae_formula or "Poids_PIB" in icae_formula,
                  f"O{ip_last} = {icae_formula[:60]}")
            # GA formula
            ga_formula = str(ws_ip.cell(ip_last, 16).value or "")  # Col P
            check("ICAE_Pays — formule GA dans dernière ligne",
                  "O" in ga_formula and "IFERROR" in ga_formula,
                  f"P{ip_last} = {ga_formula[:60]}")

        # Vérifier ICAE_Trimestriel
        if "ICAE_Trimestriel" in wb_cemac.sheetnames:
            ws_qt = wb_cemac["ICAE_Trimestriel"]
            qt_last = _actual_last_row(ws_qt)
            expected_qt_rows = 4 + len(q_cemac) - 1
            check("ICAE_Trimestriel — nombre de lignes",
                  abs(qt_last - expected_qt_rows) <= 1,
                  f"Attendu≈{expected_qt_rows}, trouvé={qt_last}")
            # Trimestre label
            qt_first = ws_qt.cell(4, 1).value
            check("ICAE_Trimestriel — premier trimestre",
                  qt_first is not None,
                  f"A4 = {qt_first}")
            # Formule AVERAGEIFS
            avg_formula = str(ws_qt.cell(4, 4).value or "")  # Col D = first country
            check("ICAE_Trimestriel — formule AVERAGEIFS",
                  "AVERAGEIFS" in avg_formula,
                  f"D4 = {avg_formula[:60]}")

        # Vérifier Poids_PIB
        if "Poids_PIB" in wb_cemac.sheetnames:
            ws_pp = wb_cemac["Poids_PIB"]
            pib_val = ws_pp.cell(4, 3).value  # C4 = PIB CMR
            check("Poids_PIB — PIB CMR",
                  pib_val is not None and float(pib_val) > 10000,
                  f"C4 = {pib_val}")
            poids_formula = str(ws_pp.cell(4, 4).value or "")
            check("Poids_PIB — formule poids",
                  "SUM" in poids_formula,
                  f"D4 = {poids_formula}")

    wb_cemac.close()
    out_cemac_path = OUTPUT_DIR / "E2E_ICAE_CEMAC.xlsx"
    out_cemac_path.write_bytes(cemac_bytes)
    print(f"  ℹ️  Export CEMAC sauvegardé : {out_cemac_path}")


# ═══════════════════════════════════════════════════════════════════════════
#  MODULE 5 — RAPPORTS : Vérification de la génération Word
# ═══════════════════════════════════════════════════════════════════════════
section("MODULE 5 — RAPPORTS (Génération Word)")

try:
    from io_utils.word_report import generate_note_nowcast, generate_note_icae
    has_word = True
except ImportError as e:
    has_word = False
    print(f"  ⚠️  Import word_report échoué : {e}")

if has_word and len(nw_results) > 0:
    # Préparer les données pour le rapport Nowcast
    perf_rows = []
    for name, r in nw_results.items():
        if not isinstance(r, dict) or "metrics" not in r:
            continue
        m = r["metrics"]
        perf_rows.append({
            "Modèle": name,
            "RMSE (in)": round(m["in_sample"].get("rmse", np.nan), 2),
            "RMSE (out)": round(m["out_sample"].get("rmse", np.nan), 2),
        })
    metrics_df = pd.DataFrame(perf_rows)
    forecasts_for_report = {m: r["forecast"] for m, r in nw_results.items()
                            if isinstance(r, dict) and "forecast" in r}

    nw_params = {
        "periodo": "T4 2022",
        "models_used": ["Bridge", "U-MIDAS", "PC", "DFM"],
        "hf_vars": ["HF1", "HF2", "HF3"],
    }

    try:
        doc_bytes = generate_note_nowcast(
            results_by_country={"SYN": {
                "metrics_df": metrics_df,
                "pib_q": pib_q,
                "forecasts": forecasts_for_report,
                "best_model": "Bridge",
                "country_name": "Synthétique",
            }},
            params=nw_params,
        )
        check("Note Nowcast Word générée", len(doc_bytes) > 1000,
              f"{len(doc_bytes)} octets")
        out_word = OUTPUT_DIR / "E2E_Note_Nowcast.docx"
        out_word.write_bytes(doc_bytes)
        print(f"  ℹ️  Note Word sauvegardée : {out_word}")
    except Exception as e:
        check("Note Nowcast Word générée", False, f"Erreur : {e}")
        traceback.print_exc()


# ═══════════════════════════════════════════════════════════════════════════
#  BILAN FINAL
# ═══════════════════════════════════════════════════════════════════════════
section("BILAN FINAL")
total = PASS + FAIL
print(f"\n  ✅ {PASS}/{total} tests passés")
print(f"  ❌ {FAIL}/{total} tests échoués")

if ERRORS:
    print(f"\n  Détails des échecs :")
    for e in ERRORS:
        print(f"    {e}")

print()
sys.exit(0 if FAIL == 0 else 1)
